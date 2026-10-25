"""Fill the downloaded official six-slide template; python-pptx is an optional authoring dependency."""
from pathlib import Path
import sys
import json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".tools/presentation-lib"))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

presentation = Presentation(ROOT / "docs/pitch/official-template.pptx")
team = json.loads((ROOT / "docs/pitch/team.json").read_text(encoding="utf-8"))
assert len(presentation.slides) == 6, "Review a changed official template before regenerating"

content = {
    0: {
        1: "PathForge AI | CYBERSTARS",
        2: "Career pathways that respect a student's skills, time, device and budget.",
        4: "Working local prototype",
        5: "12 careers · 36 skills\n33 checked learning references\nExplained matching and weekly plans\nProgress saved on the student's browser\n\nNow: three practice quizzes, initial Marathi navigation/onboarding and optional AI wording with rules fallback. Live AI verification remains pending.",
    },
    1: {
        3: "Problem hypothesis: career advice should account for practical learning constraints.",
        7: "Students may know a career goal but struggle to connect current skills with realistic next steps. A useful plan needs to reflect time, device access and budget.",
        10: "Intended users: diploma, ITI, undergraduate and first-generation job seekers in Maharashtra and similar regions. No measured beneficiary count is claimed.",
        13: "No student interviews or field pilot yet. Problem and curriculum assumptions require local review. Current evidence is a working prototype and automated checks.",
        16: "Hypothesis to test: generic plans can overlook device access, prerequisite skills and available study hours. No competitor cost or speed benchmark has been performed.",
        18: "Not measured",
        19: "Target reach is to be established through a consenting pilot.",
        21: "Aim: make the next learning step understandable and feasible without promising a job.",
    },
    2: {
        3: "Built: guest onboarding → explained matches → career gaps → weekly roadmap → local progress.",
        6: "A Next.js interface and FastAPI service calculate career matches with transparent rules. The planner schedules prerequisites, free references and activities within weekly hours, with explicit device limitations.",
        12: "Skills, interests, time, budget and device",
        18: "Weighted matching + prerequisite planning",
        24: "Three paths, exact gaps and weekly tasks",
        28: "Same SOC goal: 6h/week + laptop gives 4 illustrative weeks; 4h/week + phone gives 6 weeks and theory alternatives. This is a design mechanism, not a validated novelty or patent claim.",
    },
    3: {
        3: "Working locally; public hosting and human validation remain pending.",
        12: "✓ Working Prototype",
        19: "Next.js + FastAPI + SQLite. Deterministic core, optional Gemini wording with fallback, three practice quizzes, initial Marathi navigation/onboarding and protected expert administration. See verification for current limits.",
        22: "Current demo: zero paid API calls; free public reading references. Free hosting is conditional on provider limits, not an always-on guarantee. Pilot staffing/device costs are not yet quoted.",
        25: "Proposed: 3 months—review mappings and a small consenting pilot; 6 months—evaluate usefulness and refine; 12 months—consider expansion if evidence supports it. These are targets, not completed milestones.",
        28: "No adoption or institutional partner confirmed. Seeking educators/counsellors for curriculum review and access to consenting students for a pilot.",
    },
    4: {
        3: "Intended impact; outcomes have not yet been measured.",
        6: "Pilot planned",
        8: "5–10 consenting students proposed; not recruited",
        10: "Not measured",
        12: "No baseline time/cost comparison yet",
        14: "Not measured",
        16: "No placement or job-creation claim",
        19: "Aim: clearer career exploration for students with different devices, study hours and budgets. Phone alternatives support theory; practical work still needs appropriate access.",
        22: "Current prototype avoids paid AI. Longer-term hosting, maintenance and curriculum review need a sustainable sponsor/institution model; no revenue or funding is claimed.",
        25: "Start with local curriculum review. Track explanation comprehension, usefulness (1–5), task completion and resource availability, then expand only after reviewing evidence.",
    },
    5: {
        3: "CYBERSTARS | Aarush Tendulkar, Team Lead | Team size: 1. Guidance experts are separate.",
        6: "No photo supplied",
        7: team["applicant"]["name"],
        8: team["applicant"]["role"] + " | CYBERSTARS",
        9: team["applicant"]["qualification"] or "Qualification: to be confirmed",
        28: "Mentor / institutional partner: none confirmed",
        32: "Requested: mentorship and curriculum review",
        34: "Requested: access to a consenting local pilot",
        36: "Funding needs: not quoted; define after pilot design",
        38: "Incubation: optional future support",
        40: "Requested: accessibility and evaluation guidance",
        41: "Support would help test usefulness and feasibility with students before making impact or readiness claims.",
        42: "Contact: " + (team["applicant"]["contact"] or "") + " | Public demo URL: pending deployment",
    },
}

def replace(shape, text):
    old_size = next((run.font.size for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size), Pt(14))
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(text.split("\n")):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        paragraph.text = line
        for run in paragraph.runs:
            run.font.size = min(old_size, Pt(12 if len(text) > 120 else 18))
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor.from_string("173E38")

for slide_index, replacements in content.items():
    slide = presentation.slides[slide_index]
    for shape_index, text in replacements.items():
        replace(slide.shapes[shape_index], text)

# Preserve the template's dark-slide contrast instead of applying body ink on navy.
for slide_index, indexes in [(0, [1, 2, 4, 5]), (1, [18, 19, 21]), (5, [3, 6, 7, 8, 9, 28, 42])]:
    for index in indexes:
        for paragraph in presentation.slides[slide_index].shapes[index].text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
for run in presentation.slides[0].shapes[1].text_frame.paragraphs[0].runs:
    run.font.size = Pt(28)

table = presentation.slides[2].shapes[26].table
values = [
    ["", "Baseline to validate", "Current prototype"],
    ["Cost", "Not benchmarked", "No paid API calls"],
    ["Speed / time", "Not benchmarked", "Weekly hours constrained"],
    ["Reach / access", "Not benchmarked", "Guest; phone limits explicit"],
    ["Ease of use", "Not benchmarked", "Explained scores; pilot planned"],
]
for row, entries in zip(table.rows, values):
    for cell, text in zip(row.cells, entries):
        cell.text = text
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(10)
for column in [0, 1]:
    for paragraph in table.cell(0, column).text_frame.paragraphs:
        paragraph.font.color.rgb = RGBColor(255, 255, 255)

# Guidance experts are application data, never automatic project team members.
slide = presentation.slides[5]
original = list(slide.shapes)
for index in range(3):
    for shape in original[10 + index * 6:16 + index * 6]:
        shape._element.getparent().remove(shape._element)
box = slide.shapes.add_textbox(Inches(.5), Inches(2.9), Inches(4.6), Inches(1.1))
replace(box, "Project team: Aarush Tendulkar, Team Lead. Team size: 1.\nStudent guidance experts are managed separately in the application. No experts or institutional partners are confirmed.")
for paragraph in box.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
for index, top, height, size in [(7, 1.48, .43, 11), (8, 1.94, .3, 9), (9, 2.29, .2, 8)]:
    shape = original[index]
    shape.top, shape.height = Inches(top), Inches(height)
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)

slide = presentation.slides[0]
box = slide.shapes.add_textbox(Inches(.5), Inches(3), Inches(5.6), Inches(1.8))
replace(box, "General | Education & Employability\nWorking prototype: deterministic core + Phase 3\nAarush Tendulkar | CYBERSTARS\nLocal demonstration; public deployment pending")
for paragraph in box.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
image = ROOT / "docs/media/screenshots/roadmap-desktop.png"
if not image.exists():
    image = ROOT / "frontend/test-results/review/roadmap-desktop.png"
if image.exists():
    slide.shapes.add_picture(str(image), Inches(6.6), Inches(3.9), width=Inches(2.7))

notes = [
    "Scope: deterministic core plus optional Gemini wording with fallback, three backend-graded practice quizzes and initial Marathi navigation/onboarding. No live-provider success, complete localization, credential or pilot outcome is claimed.",
    "Problem is a hypothesis from the project brief. Do not substitute fabricated population/survey/placement statistics.",
    "Architecture: Next.js → FastAPI → pure matching/planning; SQLite stores public catalogs; localStorage stores guest progress. Source and scoring disclosures are in docs/architecture.md and docs/data-sources.md.",
    "Costs exclude engineering time, devices, connectivity, supervision and support. Vercel Hobby and Render Free have usage/cold-start limits. Public deployment has not been verified.",
    "Pilot measures: consent/date/sample size, completion rate, time to roadmap, comprehension, usefulness, task completion and working-link rate. All are planned, not collected.",
    "Applicant: Aarush Tendulkar; Team Lead, CYBERSTARS; team size 1. General Category, Education and Employability. Contact remains null; verify mandatory contact and qualification fields directly with the applicant. Guidance experts belong to the database, not the project roster. Organizer page asks PPTX while downloaded template asks PDF ≤10MB; follow the actual portal upload control. Source: https://www.sevainnovationchallenge.in/guidelines.html checked 25 Oct 2026.",
]
for slide, note in zip(presentation.slides, notes):
    slide.notes_slide.notes_text_frame.text = note
target = ROOT / "docs/pitch/PathForge_AI_CYBERSTARS_PHASE3_DRAFT.pptx"
presentation.save(target)
print(f"Created {target.name}: {len(presentation.slides)} official-template slides; confirmed applicant identity and null contact retained.")
